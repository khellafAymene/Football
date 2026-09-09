from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

doc = Document()

# Header
title = doc.add_paragraph()
title.alignment = WD_ALIGN_PARAGRAPH.RIGHT
run = title.add_run("الأتقان للتجارة والمقاولات\nAl-Atqan Trading And Contracting")
run.font.size = Pt(16)
run.bold = True

inv_title = doc.add_paragraph()
inv_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
inv_run = inv_title.add_run("INVOICE")
inv_run.font.size = Pt(14)
inv_run.bold = True

# Ref and Date Table
t_meta = doc.add_table(rows=1, cols=2)
t_meta.alignment = WD_TABLE_ALIGNMENT.CENTER
t_meta.rows[0].cells[0].text = "INVOICE REF: #002"
t_meta.rows[0].cells[1].text = "DATE: 17/08/2026"

doc.add_paragraph()

# Bill To
p_bill = doc.add_paragraph()
p_bill.add_run("BILL TO:\n").bold = True
p_bill.add_run("COMBINED GROUP ASPHALT PLANT & ROADS CONSTRUCTION\nBUZWA\nAttn.: The Finance Manager")

doc.add_paragraph()

# Main Items Table
table = doc.add_table(rows=5, cols=5)
table.alignment = WD_TABLE_ALIGNMENT.CENTER

headers = ["DESCRIPTION", "UNIT", "QTY", "RATE", "AMOUNT (QR)"]
for i, h in enumerate(headers):
    table.rows[0].cells[i].text = h
    table.rows[0].cells[i].paragraphs[0].runs[0].font.bold = True

items = [
    ("Transportation of Asphalt Material\nنقل اسفلت إلى بركة العوامر", "TON", "575.00", "9.00", "5,175.00"),
    ("Transportation of Asphalt Material\nنقل اسفلت إلى الكرعانة", "TON", "576.00", "10.00", "5,760.00"),
    ("TOTAL AMOUNT QR:", "", "", "", "10,935.00"),
    ("GRAND TOTAL (THIS INVOICE) QR:", "", "", "", "10,935.00"),
]

for row_idx, data in enumerate(items, start=1):
    for col_idx, val in enumerate(data):
        cell = table.rows[row_idx].cells[col_idx]
        cell.text = val

doc.add_paragraph()
p_words = doc.add_paragraph()
p_words.add_run("Qatari Riyals: Ten Thousand Nine Hundred Thirty-Five Riyals only").bold = True

# Footer / Signatures
p_sign = doc.add_paragraph()
p_sign.add_run("AL Atqan Trading & Contracting\nمدير العمليات: خليل\nالهاتف: 77903738")

p_contact = doc.add_paragraph()
p_contact.alignment = WD_ALIGN_PARAGRAPH.CENTER
p_contact.add_run("Tel: 44440274 / 44411985 | Fax: 44805327 | P.O.Box: 11936 Qatar - Doha | E-Mail: al.atqancontracting@gmail.com")

doc.save("Invoice_002.docx")
print("File saved successfully as Invoice_002.docx")